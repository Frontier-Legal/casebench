"""The spreadsheet a lawyer fills, one row per case, and the rows read back (decision 0049)."""

import csv
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from casebench.plugins import registry
from casebench.plugins.kinds import kind_named
from casebench.tasks.interface import TaskType

ID = "id"
PRACTICE_AREA = "practice area"
HOW = "how to fill"
# What each column takes, for the sheet that explains them.
COLUMNS = {
    ID: "Optional. The case's id; without one, the file's name and the row's number.",
    PRACTICE_AREA: "Optional. The area of law, as your team names it, such as labour or tax. "
    "The report gives each metric by practice area.",
    "question": "The question, as your agent's users would ask it.",
    "facts": "The facts the answer may rely on.",
    "constraints": "Limits on the research, such as courts or dates.",
    "instruction": "What to do with the facts or the document.",
    "document": "The path to a .txt, .md or .docx file, relative to this spreadsheet.",
    "fields": "One field per line, as name: what to take, such as party_a: the first party.",
    "reference answer": "The answer you would give. It is shown to the judge as your evidence.",
    "key points": "One per line, each with its weight, such as +5 Concludes the dismissal is "
    "invalid or -4 Relies on the repealed article. A line without a weight weighs 1.",
    "statutes": "One per line, as you would cite it.",
    "precedents": "One decision per line, as you would cite it.",
    "issues": "One per line, as clause number: what is wrong there, such as 9.2: Uncapped "
    "indemnity.",
    "expected": "One field per line, as name: value, such as party_a: Acme A.Ş.",
    "outcome": "The court's ruling, such as upheld or reversed.",
}
WIDE = 48


class SheetError(ValueError):
    """A row or file the import cannot read; the message says where."""


@dataclass(frozen=True)
class Row:
    task: str
    number: int
    cells: dict[str, str]

    def where(self, source: Path) -> str:
        return f"{source.name}, sheet {self.task}, row {self.number}"


def task_type(name: str) -> TaskType:
    return registry.load(kind_named("task"), name)


def task_names() -> list[str]:
    return [item.name for item in registry.available(kind_named("task"))]


def columns(task: TaskType) -> list[str]:
    return [ID, *task.input_fields, *task.reference_fields, PRACTICE_AREA]


def write_template(path: Path, task: str) -> None:
    """An XLSX workbook with a sheet named after the task and one that explains its columns;
    a CSV with the header row alone."""
    names = columns(task_type(task))
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8", newline="") as stream:
            csv.writer(stream).writerow(names)
        return
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = task
    sheet.append(names)
    for number, cell in enumerate(sheet[1], start=1):
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions[get_column_letter(number)].width = WIDE
    sheet.freeze_panes = "A2"
    how = book.create_sheet(HOW)
    how.append(["column", "what it takes"])
    for name in names:
        how.append([name, COLUMNS.get(name, "")])
    how.column_dimensions["A"].width = 20
    how.column_dimensions["B"].width = 100
    book.save(path)


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows(task: str, table: list[list[object]]) -> list[Row]:
    if not table:
        return []
    header = [_text(name).casefold() for name in table[0]]
    rows = []
    for number, values in enumerate(table[1:], start=2):
        cells = {name: _text(value) for name, value in zip(header, values, strict=False) if name}
        if any(cells.values()):
            rows.append(Row(task, number, cells))
    return rows


def read_rows(path: Path, task: str | None = None) -> list[Row]:
    """Every filled row of every sheet named after a task type, or of a CSV of `task`."""
    if not path.is_file():
        raise SheetError(f"{path} does not exist")
    if path.suffix.lower() == ".csv":
        if task is None:
            raise SheetError(f"{path.name} is a CSV, so name its task with --task")
        with path.open(encoding="utf-8-sig", newline="") as stream:
            return _rows(task, [list(row) for row in csv.reader(stream)])
    known = task_names()
    book = load_workbook(path, read_only=True, data_only=True)
    sheets = [sheet for sheet in book.worksheets if sheet.title in known]
    rows = []
    for sheet in sheets:
        if task is None or sheet.title == task:
            rows += _rows(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
    book.close()
    if not sheets:
        raise SheetError(
            f"{path.name} has no sheet named after a task type ({', '.join(known)}); "
            "casebench template lawyer writes one"
        )
    return rows


WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_text(path: Path) -> str:
    """The paragraphs of a Word document, one per line."""
    with zipfile.ZipFile(path) as archive:
        body = ElementTree.fromstring(archive.read("word/document.xml"))
    lines = []
    for paragraph in body.iter(f"{WORD}p"):
        parts = []
        for node in paragraph.iter():
            if node.tag == f"{WORD}t":
                parts.append(node.text or "")
            elif node.tag == f"{WORD}tab":
                parts.append("\t")
        lines.append("".join(parts))
    return "\n".join(lines)


def read_document(path: Path) -> str:
    if not path.is_file():
        raise SheetError(f"{path} does not exist")
    if path.suffix.lower() == ".docx":
        return _docx_text(path)
    return path.read_text(encoding="utf-8")
