"""The partner report as a Word document (decision 0054)."""

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.table import Table

from casebench.report.partner import (
    READING,
    PartnerReport,
    approval_lines,
    area_name,
    interval,
    percent,
)

MUTED = RGBColor(0x71, 0x71, 0x7A)


def _table(document, header: list[str], rows: list[list[str]], numeric: tuple[int, ...]) -> Table:
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
    for values in rows:
        cells = table.add_row().cells
        for n, (cell, text) in enumerate(zip(cells, values, strict=True)):
            cell.text = text
            if n in numeric:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return table


def _labelled(document, label: str, text: str) -> None:
    heading = document.add_paragraph()
    run = heading.add_run(label.upper())
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    heading.paragraph_format.space_after = Pt(0)
    document.add_paragraph(text)


def partner_docx(report: PartnerReport) -> bytes:
    document = Document()
    style = document.styles["Normal"]
    assert isinstance(style, ParagraphStyle)
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    document.add_heading(report.suite_id, level=0)
    intro = document.add_paragraph(
        f"How the AI under test, {report.target}, answered the suite's cases: run "
        f"{report.run_id}, scored on {report.scored_on}, exported on {report.exported_on}. "
        f"{report.approved} of {report.cases} cases are approved by a lawyer."
    )
    intro.runs[0].font.color.rgb = MUTED

    document.add_heading("How to read it", level=1)
    for line in READING:
        document.add_paragraph(line)
    if report.left_out:
        document.add_paragraph(f"{report.left_out} answers to unreviewed cases are left out.")

    document.add_heading("Errors", level=1)
    if report.error_rates:
        _table(
            document,
            ["Error", "What it means", "Rate", "95% interval", "Answers"],
            [
                [
                    rate.error.capitalize(),
                    f"The answer {rate.meaning}.",
                    percent(rate.rate),
                    interval(rate.low, rate.high),
                    f"{len(rate.failing)} of {rate.n}",
                ]
                for rate in report.error_rates
            ],
            numeric=(2, 3, 4),
        )
    else:
        document.add_paragraph("No metric in this run measures these errors.")

    document.add_heading("Results by task and practice area", level=1)
    _table(
        document,
        ["Task", "Measure", "Practice area", "Cases", "Score", "95% interval"],
        [
            [
                row.task_type,
                row.name,
                area_name(row.practice_area),
                str(row.n),
                percent(row.mean),
                interval(row.low, row.high),
            ]
            for row in report.metrics
        ],
        numeric=(3, 4, 5),
    )

    document.add_heading("Worked examples", level=1)
    if not report.examples:
        document.add_paragraph("No answer had one of these errors.")
    for example in report.examples:
        document.add_heading(
            f"{example.error.capitalize()} · case {example.case_id} · "
            f"{area_name(example.practice_area)}",
            level=2,
        )
        _labelled(document, "Question", example.question)
        _labelled(document, "The AI's answer", example.answer)
        for shown in example.evidence:
            _labelled(document, shown.label, shown.text)

    document.add_heading("Who approved the answer key", level=1)
    for line in approval_lines(report):
        document.add_paragraph(line)
    footer = document.add_paragraph(
        f"Suite {report.suite_id}, hash {report.suite_hash}. Run {report.run_id}. "
        f"Casebench {report.casebench_version}."
    )
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = MUTED

    out = BytesIO()
    document.save(out)
    return out.getvalue()
