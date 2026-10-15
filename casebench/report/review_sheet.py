"""A run's criteria as a spreadsheet a lawyer reviews offline, and the verdicts read back.

Each filled row becomes a lawyer's verdict in the run's reviews.jsonl, as one given in the
workbench (decision 0055).
"""

from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from casebench.core.case import Case, Fact
from casebench.core.score import CriterionVerdict, Score
from casebench.corpus.binding import CorpusConfig
from casebench.corpus.ids import SEPARATOR
from casebench.report.calibration import (
    CRITERIA,
    LawyerVerdict,
    Said,
    add_review,
    criterion_verdicts,
    read_reviews,
    said,
)
from casebench.runs.log import RunLog

REVIEW, ABOUT = "review", "about"
MARK = "casebench review sheet"
CASE, REPEAT, CRITERION = "case", "repeat", "criterion"
VERDICT, NOTE, BY = "verdict", "note", "reviewed by"
# The columns, each with its width; the last three are the lawyer's.
COLUMNS = [
    (CASE, 14),
    (REPEAT, 8),
    (CRITERION, 12),
    ("question", 40),
    ("the AI's answer", 60),
    ("what the answer must do", 40),
    ("the judge said", 12),
    ("the judge relied on", 30),
    ("evidence", 40),
    (VERDICT, 10),
    (NOTE, 40),
    (BY, 18),
]
VERDICTS = ("pass", "fail")


class ReviewSheetError(ValueError):
    """A sheet the import cannot use; the message names the row and the column."""


@dataclass(frozen=True)
class Given:
    number: int
    case_id: str
    repeat: int
    criterion_id: str
    verdict: Said
    note: str
    by: str


def _link(fact: Fact, corpora: dict[str, CorpusConfig]) -> str | None:
    """The provider's page for the decision a fact comes from, when the corpus gives `link`."""
    corpus, _, document = fact.evidence.doc.partition(SEPARATOR)
    config = corpora.get(corpus)
    if not document or config is None or not config.link:
        return None
    return config.link.replace("{id}", quote(document, safe=""))


def _evidence(case: Case | None, criterion_id: str, corpora: dict[str, CorpusConfig]):
    """What a criterion rests on: a lawyer's own words, or a court fact with its document id,
    never the court's passage. Returns the text and the first link."""
    if case is None:
        return "", None
    criterion = next((c for c in case.reference.criteria if c.id == criterion_id), None)
    if criterion is None:
        return "", None
    facts = {f.id: f for f in case.reference.facts}
    lines, link = [], None
    for fact_id in criterion.derived_from:
        fact = facts.get(fact_id)
        if fact is None:
            continue
        if fact.evidence.quote or fact.evidence.span:
            lines.append(f"{fact.kind}: {fact.value} (court record {fact.evidence.doc})")
            link = link or _link(fact, corpora)
        else:
            lines.append(f"{fact.kind}: {fact.value}")
    return "\n".join(lines), link


def _question(case: Case | None) -> str:
    if case is None:
        return ""
    inputs = case.task.inputs
    return str(inputs.get("question") or inputs.get("instruction") or "")


def _criterion_text(case: Case | None, criterion_id: str) -> str:
    if case is None:
        return criterion_id
    found = next((c for c in case.reference.criteria if c.id == criterion_id), None)
    return found.text if found else criterion_id


def _rank(scores: list[Score]) -> int:
    """Where the judges disagreed first, then failed criteria, then the rest (0050)."""
    if any(s.disagreements for s in scores):
        return 0
    if any(not v.passed for s in scores if s.metric == CRITERIA for v in s.criteria):
        return 1
    return 2


def write_review_sheet(
    log: RunLog,
    cases: list[Case],
    path: Path,
    corpora: dict[str, CorpusConfig] | None = None,
    everything: bool = False,
) -> int:
    """The rows written. Without `everything`, criteria a lawyer has given a verdict on are
    left out; with it, they are in, the lawyer's verdict filled in."""
    sets = log.score_sets()
    if not sets:
        raise ReviewSheetError(f"{log.folder} has no scores yet; score it first")
    scores = RunLog.read_scores(sets[-1])
    reviews = {(r.case_id, r.repeat, r.criterion_id): r for r in read_reviews(log.folder)}
    results = {(r.case_id, r.repeat): r for r in log.results()}
    by_id = {case.id: case for case in cases}
    grouped: dict[tuple[str, int], list[Score]] = defaultdict(list)
    for score in scores:
        grouped[(score.case_id, score.repeat)].append(score)
    order = sorted(grouped, key=lambda key: (_rank(grouped[key]), key))

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = REVIEW
    sheet.append([name for name, _ in COLUMNS])
    for number, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(number)].width = width
        sheet.cell(row=1, column=number).font = Font(bold=True)
    written = 0
    for key in order:
        case = by_id.get(key[0])
        result = results.get(key)
        answer = (result.output or result.error or "") if result else ""
        verdicts: list[CriterionVerdict] = [
            v for s in grouped[key] if s.metric == CRITERIA for v in s.criteria
        ]
        for verdict in verdicts:
            review = reviews.get((*key, verdict.criterion_id))
            if review is not None and not everything:
                continue
            evidence, link = _evidence(case, verdict.criterion_id, corpora or {})
            sheet.append(
                [
                    key[0],
                    key[1],
                    verdict.criterion_id,
                    _question(case),
                    answer,
                    _criterion_text(case, verdict.criterion_id),
                    said(verdict),
                    verdict.quote or "",
                    evidence,
                    review.verdict if review else None,
                    review.note if review else None,
                    review.by if review else None,
                ]
            )
            row = sheet.max_row
            if link:
                cell = sheet.cell(row=row, column=9)
                assert isinstance(cell, Cell)
                cell.hyperlink = link
                cell.font = Font(color="0563C1", underline="single")
            written += 1
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    allowed = DataValidation(type="list", formula1=f'"{",".join(VERDICTS)}"', allow_blank=True)
    allowed.error = "Write pass or fail."
    sheet.add_data_validation(allowed)
    allowed.add(f"J2:J{max(sheet.max_row, 2)}")
    sheet.freeze_panes = "D2"

    about = book.create_sheet(ABOUT)
    about.append([MARK])
    about.append(["run", log.info.run_id])
    about.append(["score set", sets[-1].name])
    about.append(["made", datetime.now(UTC).isoformat(timespec="seconds")])
    about.sheet_state = "hidden"
    book.save(path)
    return written


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def run_of(path: Path) -> str | None:
    """The run a review sheet was made from, or None for any other workbook."""
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        return None
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if ABOUT not in book.sheetnames:
            return None
        rows = [list(row) for row in book[ABOUT].iter_rows(values_only=True)]
    finally:
        book.close()
    if not rows or _text(rows[0][0]) != MARK:
        return None
    fields = {_text(row[0]): _text(row[1]) for row in rows[1:] if len(row) > 1}
    return fields.get("run") or None


def read_given(path: Path, by: str | None = None) -> list[Given]:
    """Every row with a verdict or a note, checked; the first bad row stops the reading."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        table = [list(row) for row in book[REVIEW].iter_rows(values_only=True)]
    finally:
        book.close()
    if not table:
        return []
    header = [_text(name).casefold() for name in table[0]]
    missing = [name for name in (CASE, REPEAT, CRITERION, VERDICT, NOTE, BY) if name not in header]
    if missing:
        raise ReviewSheetError(f"{path.name} has no column {', '.join(missing)}")
    given = []
    for number, values in enumerate(table[1:], start=2):
        cells = {name: _text(value) for name, value in zip(header, values, strict=False)}
        verdict, note = cells.get(VERDICT, "").casefold(), cells.get(NOTE, "")
        if not verdict and not note:
            continue
        where = f"{path.name}, row {number}"
        if verdict not in VERDICTS:
            shown = cells.get(VERDICT) or "nothing"
            raise ReviewSheetError(
                f"{where}, column {VERDICT}: {shown!r} is not pass or fail"
                if verdict
                else f"{where}, column {VERDICT}: the row has a note but no verdict"
            )
        said_by: Said = "pass" if verdict == "pass" else "fail"
        reviewer = cells.get(BY) or (by or "").strip()
        if not reviewer:
            raise ReviewSheetError(f"{where}, column {BY}: say who reviewed it, or pass --by")
        if not cells.get(REPEAT, "").isdigit():
            raise ReviewSheetError(f"{where}, column {REPEAT}: {cells.get(REPEAT)!r} is no number")
        given.append(
            Given(
                number=number,
                case_id=cells.get(CASE, ""),
                repeat=int(cells[REPEAT]),
                criterion_id=cells.get(CRITERION, ""),
                verdict=said_by,
                note=note,
                by=reviewer,
            )
        )
    return given


def import_reviews(log: RunLog, path: Path, given: list[Given]) -> list[LawyerVerdict]:
    """Every row is checked against the run before any is written: the criterion must still
    be in its latest score set, and overruling the judge needs a note."""
    sets = log.score_sets()
    if not sets:
        raise ReviewSheetError(f"{log.folder} has no scores yet; score it first")
    judged = criterion_verdicts(RunLog.read_scores(sets[-1]))
    made = []
    for row in given:
        where = f"{path.name}, row {row.number}"
        found = judged.get((row.case_id, row.repeat, row.criterion_id))
        if found is None:
            raise ReviewSheetError(
                f"{where}, column {CRITERION}: run {log.info.run_id} has no verdict on criterion "
                f"{row.criterion_id} of case {row.case_id}, repeat {row.repeat}"
            )
        judge = said(found)
        if row.verdict != judge and not row.note:
            raise ReviewSheetError(
                f"{where}, column {NOTE}: the judge said {judge}; overruling it needs a note"
            )
        made.append(
            LawyerVerdict(
                case_id=row.case_id,
                repeat=row.repeat,
                criterion_id=row.criterion_id,
                verdict=row.verdict,
                judge=judge,
                note=row.note or None,
                by=row.by,
            )
        )
    for verdict in made:
        add_review(log.folder, verdict)
    return made
