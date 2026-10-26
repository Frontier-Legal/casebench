"""Lawyers write cases in a spreadsheet, and research, drafting, editing and extraction tasks are
graded (decision 0049)."""

import csv
import json
import time
import zipfile
from pathlib import Path
from typing import Any

import pytest
import yaml
from click.testing import CliRunner
from openpyxl import Workbook, load_workbook

from casebench.cases.golden import new_case
from casebench.cli import main
from casebench.core.case import ReferenceSource
from casebench.suites.freeze import load
from tests.cases.agents import DRAFT
from tests.corpus.document_api import WORKS
from tests.judge.fake_judge import start_fake_judge
from tests.metrics.authorities import start_resolving_judge, write_corpus

AGENTS = Path(__file__).parent / "agents.py"
CONTRACT = """Services Agreement

1. Parties. Acme A.Ş. and Beta Ltd.
4. Term. The term is two years.
9. Liability.
9.1 Each party is liable for its own staff.
9.2 The supplier indemnifies the client without limit.
"""


def fill(path: Path, sheet: str, rows: list[dict[str, str]]) -> None:
    """Fill a template the way a lawyer would: under its header, one case per row."""
    book = load_workbook(path)
    table = book[sheet]
    header = [cell.value for cell in table[1]]
    for row in rows:
        table.append([row.get(str(name), "") for name in header])
    book.save(path)


def workbook(path: Path, sheet: str, rows: list[dict[str, str]]) -> Path:
    """A spreadsheet with one sheet, as `casebench template lawyer` writes it."""
    made = CliRunner().invoke(main, ["template", "lawyer", "--task", sheet, "--out", str(path)])
    assert made.exit_code == 0, made.output
    fill(path, sheet, rows)
    return path


def judge(port: int) -> dict[str, Any]:
    options = {"base_url": f"http://127.0.0.1:{port}", "max_tokens": 1024}
    return {"provider": "anthropic", "model": "judge-model", "options": options}


def config(tmp_path: Path, **blocks: Any) -> Path:
    path = tmp_path / "casebench.yaml"
    path.write_text(yaml.safe_dump(blocks), encoding="utf-8")
    return path


def import_(tmp_path: Path, sheet: Path, settings: Path | None = None):
    args = ["import", str(sheet), "--suite", str(tmp_path / "suite"), "--by", "Ayşe"]
    args += ["--config", str(settings or tmp_path / "none.yaml")]
    return CliRunner().invoke(main, args)


def run(tmp_path: Path, agent: str, settings: Path | None = None):
    return CliRunner().invoke(
        main,
        [
            "run",
            "--suite",
            str(tmp_path / "suite"),
            "--config",
            str(settings or tmp_path / "none.yaml"),
            "--python",
            f"{AGENTS}:{agent}",
            "--runs-dir",
            str(tmp_path / "runs"),
            "--run-id",
            "r",
            "--no-capture",
        ],
    )


def scores(tmp_path: Path) -> dict[tuple[str, str], dict[str, Any]]:
    report = json.loads((tmp_path / "runs" / "r" / "report.json").read_text(encoding="utf-8"))
    return {(s["case_id"], s["metric"]): s for s in report["scores"]}


@pytest.fixture
def key(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")


def test_template_round_trip(tmp_path: Path, key: None):
    sheet = workbook(
        tmp_path / "cases.xlsx",
        "qa",
        [
            {
                "question": "Was the appeal upheld?",
                "facts": "The lower court dismissed the claim.",
                "reference answer": "Yes, the chamber upheld the decision.",
                "key points": "+5 Mentions upheld\n-4 Mentions reversed",
                "statutes": "4857 sayılı İş Kanunu m. 18",
            },
            {"id": "bare-1", "question": "What now?", "reference answer": "Appeal in two weeks."},
        ],
    )
    template = load_workbook(sheet)
    server = start_fake_judge()

    imported = import_(tmp_path, sheet)
    ran = run(tmp_path, "answers", config(tmp_path, judge=judge(server.server_port)))
    server.shutdown()

    assert template.sheetnames == ["qa", "how to fill"]
    assert imported.exit_code == 0, imported.output
    _, cases = load(tmp_path / "suite")
    first, bare = cases
    assert first.id == "cases-2" and first.task.type == "qa"
    assert first.reference.source is ReferenceSource.LAWYER and first.reference.author == "Ayşe"
    assert [(c.text, c.weight) for c in first.reference.criteria] == [
        ("Mentions upheld", 5),
        ("Mentions reversed", -4),
    ]
    facts = {f.id: f for f in first.reference.facts}
    for criterion in first.reference.criteria:
        point = facts[criterion.derived_from[0]]
        assert point.kind == "key point" and point.value == criterion.text
        assert (point.evidence.doc, point.evidence.anchor) == ("cases.xlsx", "row 2")
        assert "answer-1" in criterion.derived_from
    assert facts["statute-1"].value == "4857 sayılı İş Kanunu m. 18"
    # A row with no key points makes no criterion; the import and the report name it.
    assert bare.reference.criteria == []
    assert "bare-1 (row 3) is unevidenced for criteria" in imported.output
    assert ran.exit_code == 0, ran.output
    report = json.loads((tmp_path / "runs" / "r" / "report.json").read_text(encoding="utf-8"))
    assert report["unevidenced"] == ["bare-1"]
    assert scores(tmp_path)[("cases-2", "criteria")]["value"] == pytest.approx(1.0)


def test_the_template_asks_for_the_question_as_the_agents_users_would_ask_it(tmp_path: Path):
    template = load_workbook(workbook(tmp_path / "cases.xlsx", "qa", []))

    how = {row[0]: row[1] for row in template["how to fill"].iter_rows(values_only=True)}

    assert how["question"] == "The question, as your agent's users would ask it."


def test_research_task(tmp_path: Path, key: None):
    server = start_resolving_judge()
    corpus = write_corpus(tmp_path / "corpus")
    settings = config(
        tmp_path,
        judge=judge(server.server_port),
        corpus={"connector": "folder", "options": {"path": corpus.as_posix()}},
    )
    sheet = workbook(
        tmp_path / "research.xlsx",
        "research",
        [
            {
                "question": "When is a dismissal without a valid reason invalid?",
                "constraints": "Yargıtay only, after 2015.",
                "statutes": "4857 sayılı İş Kanunu m. 18\n1475 sayılı Kanun m. 14",
                "precedents": "Yargıtay 9. HD E. 2019/1 K. 2020/2",
            }
        ],
    )

    imported = import_(tmp_path, sheet, settings)
    ran = run(tmp_path, "ranks", settings)
    server.shutdown()

    assert imported.exit_code == 0, imported.output
    assert ran.exit_code == 0, ran.output
    found = scores(tmp_path)
    recall, ndcg = (
        found[("research-2", "authority_recall")],
        found[("research-2", "authority_ndcg")],
    )
    # The agent ranks a wrong statute first, then two of the three gold authorities.
    assert recall["value"] == pytest.approx(2 / 3)
    gain = 1 / 1.584962500721156 + 1 / 2
    assert ndcg["value"] == pytest.approx(gain / (1 + gain))
    for score in (recall, ndcg):
        assert score["label"] == "missed key authority: 1475 sayılı Kanun m. 14"
    assert "authority_recall" in ran.output and "authority_ndcg" in ran.output


def write_corpus_of_decisions(folder: Path) -> Path:
    folder.mkdir()
    for work_id, work in WORKS.items():
        document = {"id": work_id, "jurisdiction": "TR", "text": work["text"]}
        (folder / f"{work_id}.json").write_text(json.dumps(document), encoding="utf-8")
    return folder


def test_drafting_graded_by_evidenced_criteria(tmp_path: Path, key: None):
    server = start_fake_judge()
    corpus = write_corpus_of_decisions(tmp_path / "decisions")
    settings = config(
        tmp_path,
        judge=judge(server.server_port),
        corpus={"connector": "folder", "options": {"path": corpus.as_posix()}},
    )
    sheet = workbook(
        tmp_path / "drafts.xlsx",
        "drafting",
        [
            {
                "instruction": "Draft the claim's argument on the dismissal.",
                "facts": "The employee was dismissed by text message.",
                "key points": "+5 Mentions invalid\n+3 Mentions reinstatement\n-4 Mentions waived",
            }
        ],
    )

    # No corpus at import: the drafting task itself has every citation checked.
    imported = import_(tmp_path, sheet)
    ran = run(tmp_path, "drafts", settings)
    server.shutdown()

    assert imported.exit_code == 0, imported.output
    assert ran.exit_code == 0, ran.output
    found = scores(tmp_path)
    criteria = found[("drafts-2", "criteria")]
    verdicts = {v["criterion_id"]: v for v in criteria["criteria"]}
    met = {cid: (v["weight"], v["met"]) for cid, v in verdicts.items()}
    assert met == {"k1": (5, True), "k2": (3, False), "k3": (-4, True)}
    # Each verdict that holds rests on a quote of the draft.
    for cid in ("k1", "k3"):
        assert verdicts[cid]["quote"] and verdicts[cid]["quote"] in DRAFT
    assert criteria["value"] == pytest.approx((5 - 4) / 8)
    assert criteria["label"] == "failed: k2, k3"
    # Every citation in the draft is checked against the corpus.
    checked = found[("drafts-2", "citation_integrity")]["citations"]
    assert {c["citation"] for c in checked} >= {"Yargıtay 9. HD, E. 2023/1234 K. 2024/567"}
    assert all(c["label"] for c in checked)


def docx(path: Path, text: str) -> Path:
    """A Word document with one paragraph per line."""
    body = "".join(f"<w:p><w:r><w:t>{line}</w:t></w:r></w:p>" for line in text.splitlines())
    xml = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}</w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", xml)
    return path


def test_editing_keeps_what_it_should_not_touch(tmp_path: Path):
    docx(tmp_path / "contract.docx", CONTRACT)
    sheet = workbook(
        tmp_path / "edits.xlsx",
        "editing",
        [
            {
                "instruction": "Fix the clauses that create unusual liability.",
                "document": "contract.docx",
                "issues": "9.2: Uncapped indemnity",
            }
        ],
    )

    imported = import_(tmp_path, sheet)
    ran = run(tmp_path, "edits")

    assert imported.exit_code == 0, imported.output
    assert ran.exit_code == 0, ran.output
    found = scores(tmp_path)
    issues, kept = found[("edits-2", "issues_found")], found[("edits-2", "edit_preservation")]
    assert (issues["value"], issues["label"]) == (1.0, "all issues found")
    assert kept["label"] == "changed outside the edit: 4"
    assert "-4. Term. The term is two years." in kept["explanation"]
    assert "+4. Term. This agreement runs for two years." in kept["explanation"]
    # The other clauses outside the issue are kept: the preamble, 1, 9 and 9.1.
    assert kept["value"] == pytest.approx(4 / 5)


def test_extraction_per_field(tmp_path: Path):
    (tmp_path / "contract.txt").write_text(CONTRACT, encoding="utf-8")
    sheet = tmp_path / "fields.csv"
    made = CliRunner().invoke(
        main, ["template", "lawyer", "--task", "extraction", "--out", str(sheet)]
    )
    assert made.exit_code == 0, made.output
    with sheet.open("a", encoding="utf-8", newline="") as stream:
        csv.writer(stream).writerow(
            [
                "",
                "contract.txt",
                "party_a: the first party\nparty_b: the second party\ndate: the date signed",
                "party_a: Acme A.Ş.\nparty_b: Beta Ltd\ndate: 2024-01-01",
            ]
        )

    unnamed = import_(tmp_path, sheet)
    imported = CliRunner().invoke(
        main,
        ["import", str(sheet), "--suite", str(tmp_path / "suite"), "--by", "Ayşe",
         "--task", "extraction", "--config", str(tmp_path / "none.yaml")],
    )  # fmt: skip
    ran = run(tmp_path, "extracts")

    assert unnamed.exit_code != 0 and "name its task with --task" in unnamed.output
    assert imported.exit_code == 0, imported.output
    assert ran.exit_code == 0, ran.output
    score = scores(tmp_path)[("fields-2", "fields")]
    assert score["value"] == pytest.approx(2 / 3)
    assert score["label"] == "wrong fields: date"
    assert "party_a: exact" in score["explanation"]
    assert "party_b: normalised" in score["explanation"]
    assert "date: expected '2024-01-01', got '2024-01-02'" in score["explanation"]


def test_a_blank_template_imports_nothing(tmp_path: Path):
    sheet = workbook(tmp_path / "empty.xlsx", "qa", [])
    other = Workbook()
    assert other.active is not None
    other.active.title = "notes"
    other.save(tmp_path / "notes.xlsx")

    empty = import_(tmp_path, sheet)
    unknown = import_(tmp_path, tmp_path / "notes.xlsx")

    assert empty.exit_code != 0 and "has no filled row" in empty.output
    assert unknown.exit_code != 0 and "no sheet named after a task type" in unknown.output


def test_a_long_line_is_read_in_one_pass():
    """A key point or an expected field padded with spaces, as a form can send it, is read at
    once and to the same values."""
    padding = " " * 200_000
    started = time.perf_counter()
    case = new_case(
        "extraction",
        {"document": "Acme A.Ş. and Beta Ltd.", "expected": f"party_a:{padding}Acme A.Ş."},
        "fields-1",
        "Ayşe",
    )
    pointed = new_case(
        "qa", {"question": "Q?", "key points": f"+5{padding}Says yes"}, "qa-1", "Ayşe"
    )

    assert time.perf_counter() - started < 2
    assert [(f.id, f.value) for f in case.reference.facts] == [("party_a", "Acme A.Ş.")]
    assert [(c.weight, c.text) for c in pointed.reference.criteria] == [(5, "Says yes")]
