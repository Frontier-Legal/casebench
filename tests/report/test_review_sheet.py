"""Lawyers review a run in a spreadsheet when there is no server (decision 0055)."""

from pathlib import Path
from typing import Any

from click.testing import CliRunner
from openpyxl import load_workbook

from casebench.cli import main
from casebench.core.case import Approval, ReferenceSource
from casebench.core.score import CriterionVerdict, Direction, Score, ScoreKind
from casebench.report.calibration import lawyers_win, read_reviews
from casebench.runs.log import Result, RunInfo, RunLog
from casebench.runs.producer import Producer
from casebench.suites.freeze import freeze
from casebench.workbench.items import queue
from casebench.workbench.sources import Workspace
from casebench.workbench.store import Store
from tests.workbench.test_review_queue import TEXT, reviewed_case, scored_run

LINK = "https://decisions.example/view?doc={id}"


def sheet_rows(path: Path) -> list[dict[str, Any]]:
    book = load_workbook(path)
    table = [list(row) for row in book["review"].iter_rows(values_only=True)]
    return [dict(zip(map(str, table[0]), row, strict=True)) for row in table[1:]]


def fill(path: Path, filled: dict[int, dict[str, str]]) -> None:
    """Writes into the given rows, numbered as in the spreadsheet."""
    book = load_workbook(path)
    sheet = book["review"]
    header = [cell.value for cell in sheet[1]]
    for row, values in filled.items():
        for name, value in values.items():
            sheet.cell(row=row, column=header.index(name) + 1, value=value)
    book.save(path)


def test_review_round_trip(tmp_path: Path):
    suite = freeze("suite", [reviewed_case()], tmp_path / "suites" / "suite")
    log = scored_run(tmp_path / "runs", reviewed_case(), suite.content_hash)
    config = tmp_path / "casebench.yaml"
    config.write_text(
        f"corpora:\n  corpus:\n    connector: folder\n    options: {{path: decisions}}\n"
        f"    link: '{LINK}'\n",
        encoding="utf-8",
    )
    sheet = tmp_path / "review-r1.xlsx"
    exported = CliRunner().invoke(
        main,
        ["export", "--review", str(log.folder), "--out", str(sheet), "--suites",
         str(tmp_path / "suites"), "--config", str(config)],
    )  # fmt: skip
    [row] = sheet_rows(sheet)
    evidence = load_workbook(sheet)["review"]["I2"]
    fill(
        sheet,
        {2: {"verdict": "Fail", "note": "m. 18 is cited for reinstatement", "reviewed by": "Ayşe"}},
    )
    store = Store(f"sqlite:///{tmp_path / 'casebench.db'}")
    imported = CliRunner().invoke(
        main,
        ["import", str(sheet), "--runs-dir", str(tmp_path / "runs"), "--database",
         f"sqlite:///{tmp_path / 'casebench.db'}"],
    )  # fmt: skip
    [review] = read_reviews(log.folder)
    scores = lawyers_win(RunLog.read_scores(log.score_sets()[-1]), read_reviews(log.folder))
    [criteria] = [s for s in scores if s.metric == "criteria"]
    workspace = Workspace(tmp_path / "runs", [tmp_path / "suites"])

    assert exported.exit_code == 0, exported.output
    assert (row["case"], row["repeat"], row["criterion"]) == ("answer-1", 0, "c-statute-1")
    assert row["question"] == "Is the dismissal valid?"
    assert row["the AI's answer"].startswith("The dismissal is invalid under İş Kanunu m. 18")
    assert row["the judge said"] == "pass" and row["the judge relied on"] == "İş Kanunu m. 18"
    # The evidence links to the decision and copies none of its text.
    assert row["evidence"] == "statute: 4857 sayılı İş Kanunu m. 18 (court record corpus:tr-1)"
    assert evidence.hyperlink is not None
    assert evidence.hyperlink.target == "https://decisions.example/view?doc=tr-1"
    sheet_text = " ".join(str(v) for r in sheet_rows(sheet) for v in r.values())
    passages = ("geçerli sebebe dayanmalıdır", "işe iadenin sonuçlarını", "kararını onamıştır")
    assert all(p in TEXT for p in passages) and not any(p in sheet_text for p in passages)
    assert imported.exit_code == 0, imported.output
    assert "1 overruling the judge" in imported.output
    # The same verdict a workbench review writes, with the lawyer's name and the time.
    assert (review.verdict, review.judge, review.by) == ("fail", "pass", "Ayşe")
    assert review.note == "m. 18 is cited for reinstatement" and review.at
    [verdict] = criteria.criteria
    assert (verdict.met, verdict.reviewed_by, verdict.judge_met) == (False, "Ayşe", True)
    assert not [i for i in queue(workspace, store, "Can") if i.kind == "answer"]
    logged = [(e.action, e.by, e.subject) for e in store.events()]
    assert ("verdict", "Ayşe", "c-statute-1") in logged
    assert ("imported", "Ayşe", "review-r1.xlsx") in logged


def many(runs: Path, answers: int) -> RunLog:
    """A run of `answers` cases, each with one criterion the judge passed."""
    log = RunLog.create(
        runs / "r1", RunInfo(run_id="r1", suite_id="suite", suite_hash="sha256:x", target="a")
    )
    scores = []
    for n in range(answers):
        log.append(Result(case_id=f"c{n:03}", output=f"Answer {n}."))
        verdict = CriterionVerdict(criterion_id="k1", weight=1, met=True, quote=f"Answer {n}")
        scores.append(
            Score(
                case_id=f"c{n:03}",
                metric="criteria",
                metric_version="1",
                kind=ScoreKind.LLM,
                direction=Direction.HIGHER,
                value=1.0,
                criteria=[verdict],
                task_type="qa",
                reference_source=ReferenceSource.LAWYER,
                approval=Approval.APPROVED,
            )
        )
    log.write_score_set(scores, Producer(metrics={}))
    return log


def test_partial_sheets_import(tmp_path: Path):
    log = many(tmp_path / "runs", 200)
    sheet = tmp_path / "review.xlsx"
    CliRunner().invoke(main, ["export", "--review", str(log.folder), "--out", str(sheet)])
    fill(sheet, {row: {"verdict": "pass"} for row in range(2, 62)})
    store = Store(f"sqlite:///{tmp_path / 'casebench.db'}")

    imported = CliRunner().invoke(
        main, ["import", str(sheet), "--runs-dir", str(tmp_path / "runs"), "--by", "Can"]
    )
    waiting = [i for i in queue(Workspace(tmp_path / "runs", []), store, "Ayşe")]

    assert len(sheet_rows(sheet)) == 200
    assert imported.exit_code == 0, imported.output
    assert len(read_reviews(log.folder)) == 60
    assert {r.by for r in read_reviews(log.folder)} == {"Can"}
    assert len(waiting) == 140
    # A second sheet leaves out what lawyers already reviewed.
    again = tmp_path / "again.xlsx"
    CliRunner().invoke(main, ["export", "--review", str(log.folder), "--out", str(again)])
    assert len(sheet_rows(again)) == 140


def test_a_bad_row_is_named(tmp_path: Path):
    log = many(tmp_path / "runs", 5)
    sheet = tmp_path / "review.xlsx"
    CliRunner().invoke(main, ["export", "--review", str(log.folder), "--out", str(sheet)])
    fill(
        sheet,
        {
            2: {"verdict": "pass", "reviewed by": "Ayşe"},
            4: {"verdict": "maybe", "reviewed by": "Ayşe"},
        },
    )

    imported = CliRunner().invoke(
        main, ["import", str(sheet), "--runs-dir", str(tmp_path / "runs")]
    )

    assert imported.exit_code != 0
    assert "review.xlsx, row 4, column verdict: 'maybe' is not pass or fail" in imported.output
    assert "nothing was imported" in imported.output
    assert read_reviews(log.folder) == []
