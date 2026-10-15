"""A run's report for partners: its error rates with their failing answers (decision 0054)."""

import csv
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import pytest
from click.testing import CliRunner
from docx import Document
from fastapi.testclient import TestClient

from casebench.cli import main
from casebench.core.case import (
    Approval,
    Case,
    Criterion,
    Fact,
    Reference,
    ReferenceSource,
    Task,
)
from casebench.core.score import (
    CitationCheck,
    CitationLabel,
    CriterionVerdict,
    Direction,
    Score,
    ScoreKind,
)
from casebench.report.calibration import record_verdict
from casebench.report.partner import partner_report
from casebench.runs.log import Result, RunInfo, RunLog
from casebench.runs.producer import Producer
from casebench.suites.freeze import freeze, load
from casebench.workbench.app import create_app
from casebench.workbench.items import answer_key
from casebench.workbench.sources import Workspace
from casebench.workbench.store import Store
from tests.workbench.test_review_queue import DECISION, HOST, as_, passage, write_corpus

AREAS = {"c1": "labour", "c2": "labour", "c3": "tax", "c4": "tax"}


def case(case_id: str) -> Case:
    """A court case asking for the ruling and the statute, approved by Ayşe."""
    facts = [
        Fact(
            id="outcome-1",
            kind="outcome",
            value="upheld",
            evidence=passage("yerel mahkeme kararını onamıştır"),
        ),
        Fact(
            id="statute-1",
            kind="statute",
            value="4857 sayılı İş Kanunu m. 18",
            evidence=passage("4857 sayılı İş Kanunu m. 18"),
        ),
    ]
    return Case(
        id=case_id,
        task=Task.model_validate({"type": "qa", "question": f"Was dismissal {case_id} valid?"}),
        reference=Reference(
            source=ReferenceSource.COURT,
            facts=facts,
            criteria=[
                Criterion(
                    id="c-outcome-1",
                    text='Passes if the answer concludes that the court ruled "upheld".',
                    weight=1,
                    derived_from=["outcome-1"],
                    kind="outcome",
                )
            ],
            author="Ayşe",
        ),
        metrics=["criteria", "citation_integrity", "outcome", "statute"],
        approval=Approval.APPROVED,
        practice_area=AREAS[case_id],
    )


def scores_for(case_id: str) -> list[Score]:
    """c1 cites a decision that does not exist and one that says something else, and misses
    the ruling; c2 relies on a statute the court did not apply; c3 states the wrong ruling."""
    common: dict[str, Any] = {
        "case_id": case_id,
        "kind": ScoreKind.LLM,
        "direction": Direction.HIGHER,
        "metric_version": "1",
        "task_type": "qa",
        "reference_source": ReferenceSource.COURT,
        "approval": Approval.APPROVED,
        "practice_area": AREAS[case_id],
    }
    verified = CitationCheck(
        citation="Yargıtay 9. HD E. 2019/1 K. 2020/2",
        claim="the dismissal needs a valid reason",
        label=CitationLabel.VERIFIED,
        document_id=DECISION,
    )
    citations = {
        "c1": [
            CitationCheck(
                citation="Yargıtay 3. HD, E. 2021/999",
                claim="see also",
                label=CitationLabel.FABRICATED,
            ),
            CitationCheck(
                citation="Yargıtay 9. HD E. 2019/1 K. 2020/2",
                claim="reinstatement is automatic",
                label=CitationLabel.MISGROUNDED,
                document_id=DECISION,
                reason="the decision says nothing on reinstatement",
            ),
        ],
    }.get(case_id, [verified])
    outcome = {"c3": "wrong outcome"}.get(case_id, "correct")
    statute = {"c2": "wrong statute"}.get(case_id, "correct")
    met = case_id != "c1"
    criterion = CriterionVerdict(
        criterion_id="c-outcome-1",
        weight=1,
        met=met,
        quote="upheld" if met else None,
        kind="outcome",
    )
    return [
        Score(metric="criteria", value=float(met), criteria=[criterion], **common),
        Score(metric="citation_integrity", value=0.5, citations=citations, **common),
        Score(metric="outcome", value=float(outcome == "correct"), label=outcome, **common),
        Score(metric="statute", value=float(statute == "correct"), label=statute, **common),
    ]


@pytest.fixture
def project(tmp_path: Path) -> dict[str, Any]:
    cases = [case(c) for c in AREAS]
    # Can approved c4; c5 waits for a lawyer and the run did not answer it.
    cases[3] = cases[3].model_copy(update={"approved_by": "Can"})
    draft = case("c1").model_copy(update={"id": "c5", "approval": Approval.DRAFT})
    suite = freeze("suite", [*cases, draft], tmp_path / "suites" / "suite")
    log = RunLog.create(
        tmp_path / "runs" / "r1",
        RunInfo(run_id="r1", suite_id="suite", suite_hash=suite.content_hash, target="agent"),
    )
    for c in cases:
        log.append(Result(case_id=c.id, output=f"The court upheld it in {c.id}."))
    log.write_score_set([s for c in AREAS for s in scores_for(c)], Producer(metrics={}))
    workspace = Workspace(
        tmp_path / "runs", [tmp_path / "suites"], write_corpus(tmp_path / "decisions")
    )
    store = Store(f"sqlite:///{tmp_path / 'casebench.db'}")
    return {"tmp": tmp_path, "workspace": workspace, "store": store, "log": log}


def client(project: dict[str, Any]) -> TestClient:
    app = create_app(project["workspace"], project["store"], [HOST])
    return TestClient(app, base_url=f"http://{HOST}")


def test_error_breakdown(project: dict[str, Any]):
    app = client(project)
    app.put("/api/me", json={"name": "Ayşe", "view": "review"})
    listed = app.get("/api/runs", headers=as_("Ayşe")).json()
    before = app.get("/api/runs/r1/report", headers=as_("Ayşe")).json()
    # A lawyer finds c1 does reach the ruling; the judge's verdict no longer counts.
    record_verdict(project["log"].folder, "c1", 0, "c-outcome-1", "pass", "Can", "it says so")
    after = app.get("/api/runs/r1/report", headers=as_("Ayşe")).json()
    opened = app.get(f"/api/items/{answer_key('r1', 'c1', 0)}", headers=as_("Ayşe")).json()

    assert listed == [{"run_id": "r1", "suite_id": "suite", "target": "agent", "scored": True}]
    rates = {r["error"]: r for r in before["rates"]}
    assert list(rates) == [
        "fabricated citation", "unsupported claim", "wrong outcome", "wrong statute",
    ]  # fmt: skip
    shown = {
        name: (r["n"], r["rate"], [f["case_id"] for f in r["failing"]]) for name, r in rates.items()
    }
    assert shown == {
        "fabricated citation": (4, 0.25, ["c1"]),
        "unsupported claim": (4, 0.25, ["c1"]),
        "wrong outcome": (4, 0.5, ["c1", "c3"]),
        "wrong statute": (4, 0.25, ["c2"]),
    }
    for rate in rates.values():
        assert rate["low"] <= rate["rate"] <= rate["high"]
    [fabricated] = rates["fabricated citation"]["failing"]
    assert fabricated["key"] == answer_key("r1", "c1", 0)
    assert fabricated["title"] == "Was dismissal c1 valid?"
    wrong = next(r for r in after["rates"] if r["error"] == "wrong outcome")
    assert [f["case_id"] for f in wrong["failing"]] == ["c3"]
    # The failing answer opens with its question, the answer and the court record.
    assert opened["inputs"][0]["text"] == "Was dismissal c1 valid?"
    assert opened["answer"] == "The court upheld it in c1."
    outcome = next(f for f in opened["reference"] if f["kind"] == "outcome")
    assert (
        outcome["quote"] == "yerel mahkeme kararını onamıştır" and outcome["document"] == DECISION
    )
    labels = [c["label"] for c in opened["citations"]]
    assert labels == ["fabricated", "misgrounded"]


def test_a_case_names_its_practice_area_and_approver(tmp_path: Path):
    """The spreadsheet's practice area column and its author reach the case; a case approved
    before the approver was recorded names its reference's author; a case without either keeps
    its bytes, so its suite keeps its hash (decision 0054)."""
    sheet = tmp_path / "cases.csv"
    with sheet.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["question", "key points", "practice area"])
        writer.writerow(["Is overtime compulsory?", "+1 Says it needs consent", "labour"])
    imported = CliRunner().invoke(
        main,
        ["import", str(sheet), "--suite", str(tmp_path / "suite"), "--by", "Ayşe", "--task",
         "qa", "--config", str(tmp_path / "none.yaml")],
    )  # fmt: skip
    _, [written] = load(tmp_path / "suite")
    older = case("c1").model_copy(update={"practice_area": None})
    first = freeze("old", [older], tmp_path / "old")
    again = freeze("old", load(tmp_path / "old")[1], tmp_path / "old")

    assert imported.exit_code == 0, imported.output
    assert (written.practice_area, written.approved_by, written.approver) == (
        "labour",
        "Ayşe",
        "Ayşe",
    )
    assert older.approved_by is None and older.approver == "Ayşe"
    assert "practice_area" not in older.model_dump(mode="json")
    assert again.content_hash == first.content_hash


def test_export_for_partners(project: dict[str, Any]):
    """Each metric by task and practice area with its interval, worked examples beside their
    evidence, and who approved the golden data, as PDF, DOCX and HTML."""
    run_dir = project["tmp"] / "runs" / "r1"
    suites = str(project["tmp"] / "suites")
    made = {
        kind: CliRunner().invoke(
            main,
            [
                "report",
                str(run_dir),
                "--export",
                str(project["tmp"] / f"report.{kind}"),
                "--suites",
                suites,
            ],
        )
        for kind in ("pdf", "docx", "html")
    }
    report = partner_report(project["log"], load(project["tmp"] / "suites" / "suite")[1])
    words = "\n".join(
        p.text for p in Document(str(project["tmp"] / "report.docx")).paragraphs
    ) + "\n".join(
        cell.text
        for table in Document(str(project["tmp"] / "report.docx")).tables
        for row in table.rows
        for cell in row.cells
    )
    pdf = (project["tmp"] / "report.pdf").read_bytes()
    downloaded = client(project).get("/api/runs/r1/export/docx", headers=as_("Ayşe"))
    unknown = client(project).get("/api/runs/r1/export/xlsx", headers=as_("Ayşe"))

    assert all(ran.exit_code == 0 for ran in made.values()), {k: r.output for k, r in made.items()}
    rows = {(r.task_type, r.name, r.practice_area): (r.n, r.mean) for r in report.metrics}
    assert rows[("qa", "Ruling correct", "labour")] == (2, 1.0)
    assert rows[("qa", "Ruling correct", "tax")] == (2, 0.5)
    assert all(r.low <= r.mean <= r.high for r in report.metrics)
    assert [a.model_dump() for a in report.approvers] == [
        {"name": "Ayşe", "cases": 3},
        {"name": "Can", "cases": 1},
    ]
    assert (report.cases, report.approved, report.unreviewed) == (5, 4, 1)
    fabricated = next(e for e in report.examples if e.error == "fabricated citation")
    assert fabricated.question == "Was dismissal c1 valid?"
    assert fabricated.answer == "The court upheld it in c1."
    assert fabricated.evidence[0].text.startswith("Yargıtay 3. HD, E. 2021/999. The corpus:")
    statute = next(e for e in report.examples if e.error == "wrong statute")
    assert statute.evidence[0].text == "4857 sayılı İş Kanunu m. 18: “4857 sayılı İş Kanunu m. 18”"
    for text in (
        "Ruling correct",
        "labour",
        "tax",
        "Ayşe approved 3 cases.",
        "Can approved 1 case.",
        "Was dismissal c1 valid?",
        "The court upheld it in c1.",
        "Fabricated citation",
    ):
        assert text in words, text
    assert pdf.startswith(b"%PDF-") and b"NotoSans" in pdf
    assert downloaded.status_code == 200
    assert downloaded.headers["content-disposition"] == (
        'attachment; filename="casebench-suite-r1.docx"'
    )
    assert unknown.status_code == 422


def test_a_static_copy_for_people_with_no_access(project: dict[str, Any], tmp_path: Path):
    """One HTML file that loads nothing from elsewhere: no script, stylesheet, font or image."""
    target = tmp_path / "partners.html"
    ran = CliRunner().invoke(
        main,
        ["report", str(project["tmp"] / "runs" / "r1"), "--export", str(target), "--suites",
         str(project["tmp"] / "suites")],
    )  # fmt: skip
    page = target.read_text(encoding="utf-8")
    tags: list[tuple[str, dict[str, str | None]]] = []

    class Tags(HTMLParser):
        def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
            tags.append((tag, dict(attrs)))

    Tags().feed(page)

    assert ran.exit_code == 0, ran.output
    assert page.startswith("<!doctype html>") and '<meta charset="utf-8">' in page
    names = {tag for tag, _ in tags}
    assert not names & {"script", "link", "img", "iframe", "object", "embed"}
    assert not any(value and "://" in value for _, attrs in tags for key, value in attrs.items())
    assert "@import" not in page and "url(" not in page
    assert "Who approved the answer key" in page and "Ayşe approved 3 cases." in page


def test_an_export_refuses_a_suite_that_changed(project: dict[str, Any]):
    """The approvers named must be those of the answer key the run was scored against."""
    folder = project["tmp"] / "suites" / "suite"
    _, cases = load(folder)
    freeze("suite", [c.model_copy(update={"practice_area": "tax"}) for c in cases], folder)

    ran = CliRunner().invoke(
        main,
        ["report", str(project["tmp"] / "runs" / "r1"), "--export",
         str(project["tmp"] / "x.pdf"), "--suites", str(project["tmp"] / "suites")],
    )  # fmt: skip

    assert ran.exit_code != 0 and "changed since run r1 ran it" in ran.output
